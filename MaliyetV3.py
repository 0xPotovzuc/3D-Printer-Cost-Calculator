import tkinter as tk
from tkinter import ttk, messagebox, simpledialog
import json
import os
import sys
import traceback
import webbrowser
from typing import Dict, List, Any, Tuple

# YENİ EKLENDİ: Excel işlemleri için gerekli kütüphane
try:
    import openpyxl
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils.exceptions import InvalidFileException
except ImportError:
    messagebox.showerror(
        "Eksik Kütüphane",
        "Bu uygulamanın Excel özelliğini kullanabilmek için 'openpyxl' kütüphanesi gereklidir.\n\nLütfen terminale 'pip install openpyxl' yazarak kurun."
    )
    sys.exit()


# --- Uygulama Sabitleri ---
APP_NAME = "3D Baskı Maliyet & Kar Hesaplayıcı Pro"
APP_VERSION = "3.0"
DATA_FILE_NAME = "app_data.json"
EXCEL_FILE_NAME = "satislar.xlsx"
# JSON veri yapısı için anahtarlar
KEY_FILAMENTS = "filament_data"
KEY_ELECTRICITY = "electricity_prices"
KEY_SHIPPING = "shipping_prices"
KEY_MARKETPLACES = "marketplaces"
KEY_SETTINGS = "settings"

# Varsayılan değerler
DEFAULT_ELECTRICITY = {"ticari": 5.54, "mesken": 3.11}
DEFAULT_MARKETPLACES = {
    "Diğer/Yok": {"rate": 0.0, "fixed": 0.0}
}
DEFAULT_SETTINGS = {
    "costing_electricity_type": "ticari",
    "default_device_power": 130.0
}
DISCORD_USER_ID = "351410962336841748"

# --- Yardımcı Sınıflar ---

# YENİ EKLENDİ: Geliştirilmiş ürün adı giriş penceresi
class CustomAskStringDialog(tk.Toplevel):
    """Daha iyi bir kullanıcı arayüzü için simpledialog'un yerini alan özel diyalog."""
    def __init__(self, parent, title, prompt):
        super().__init__(parent)
        self.transient(parent)
        self.title(title)
        self.parent = parent
        self.result = None

        self.body = ttk.Frame(self, padding="10 10 10 10")
        self.body.pack(expand=True, fill="both")

        ttk.Label(self.body, text=prompt, wraplength=250, justify="left").pack(pady=(0, 10), anchor="w")

        self.entry = ttk.Entry(self.body, width=40)
        self.entry.pack(pady=5, fill="x", expand=True)
        self.entry.focus_set()

        button_frame = ttk.Frame(self.body)
        button_frame.pack(pady=(10, 0), fill="x", expand=True)

        ok_button = ttk.Button(button_frame, text="Tamam", command=self.on_ok)
        ok_button.pack(side="right", padx=(5, 0))
        cancel_button = ttk.Button(button_frame, text="İptal", command=self.on_cancel)
        cancel_button.pack(side="right")

        self.bind("<Return>", self.on_ok)
        self.bind("<Escape>", self.on_cancel)

        # Pencereyi ortala
        self.update_idletasks()
        parent_x = parent.winfo_rootx()
        parent_y = parent.winfo_rooty()
        parent_width = parent.winfo_width()
        parent_height = parent.winfo_height()
        self_width = self.winfo_width()
        self_height = self.winfo_height()
        x = parent_x + (parent_width - self_width) // 2
        y = parent_y + (parent_height - self_height) // 2
        self.geometry(f"+{x}+{y}")

        self.protocol("WM_DELETE_WINDOW", self.on_cancel)
        self.grab_set()
        self.wait_window(self)

    def on_ok(self, event=None):
        self.result = self.entry.get().strip()
        if not self.result:
            messagebox.showwarning("Giriş Gerekli", "Ürün adı boş bırakılamaz.", parent=self)
            return
        self.destroy()

    def on_cancel(self, event=None):
        self.result = None
        self.destroy()

# DEĞİŞTİRİLDİ: ToolTip sınıfı artık kullanılmadığı için kaldırıldı.
# class ToolTip:
#     ...

class CostCalculatorApp:
    """
    3D Baskı Maliyet ve Kar Hesaplama uygulamasının ana sınıfı.
    Tüm UI bileşenlerini, veri yönetimini ve hesaplama mantığını içerir.
    """
    def __init__(self, root: tk.Tk):
        self.root = root
        self._setup_window()
        self._setup_styles()

        self.status_var = tk.StringVar()

        # --- Veri Depolama ---
        self.filaments: List[Tuple[str, float]] = []
        self.electricity_prices: Dict[str, float] = DEFAULT_ELECTRICITY.copy()
        self.shipping_prices: Dict[str, List[Tuple[str, float]]] = {}
        self.marketplaces: Dict[str, Dict[str, float]] = {}
        self.settings: Dict[str, Any] = DEFAULT_SETTINGS.copy()
        
        self.selected_marketplace_name: str | None = None

        self.last_calculation_results: Dict[str, Any] | None = None
        self.last_calculation_inputs: Dict[str, Any] | None = None

        self._load_data()
        self._create_widgets()
        self.update_status("Uygulama başarıyla başlatıldı.")

    def _setup_window(self):
        """Ana pencere ayarlarını yapar."""
        self.root.title(f"{APP_NAME} - v{APP_VERSION}")
        self.root.geometry("1100x800")
        self.root.minsize(950, 650)
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)

    def _setup_styles(self):
        """Uygulama genelinde kullanılacak ttk stillerini ayarlar."""
        self.style = ttk.Style()
        themes = self.style.theme_names()
        if 'vista' in themes:
            self.style.theme_use('vista')
        elif 'aqua' in themes:
            self.style.theme_use('aqua')
        else:
            self.style.theme_use('clam')

        self.bg_color = "#f7f7f7"
        self.text_color = "#212121"
        self.header_color = "#0d47a1"
        self.highlight_color = "#bbdefb"
        self.result_bg_color = "#ffffff"

        self.style.configure("TFrame", background=self.bg_color)
        self.style.configure("TLabel", background=self.bg_color, font=("Segoe UI", 10), foreground=self.text_color)
        self.style.configure("Header.TLabel", font=("Segoe UI", 12, "bold"), foreground=self.header_color)
        self.style.configure("TButton", font=("Segoe UI", 10, "bold"), padding=8)
        
        # DEĞİŞTİRİLDİ: Titremeye neden olan 'active' durumundaki arkaplan rengi değişikliği kaldırıldı.
        # self.style.map("TButton", background=[('active', self.highlight_color)])
        
        self.style.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"))
        self.style.configure("TLabelframe", background=self.bg_color, padding=10)
        self.style.configure("TLabelframe.Label", background=self.bg_color, foreground=self.text_color, font=("Segoe UI", 11, "bold"))
        self.style.configure("Status.TLabel", font=("Segoe UI", 9), padding=5, background="#e0e0e0")
        self.style.configure("Results.TFrame", background=self.result_bg_color)
        self.style.configure("Results.TLabel", background=self.result_bg_color)
        self.style.configure("Results.TLabelframe", background=self.result_bg_color, relief="flat")
        self.style.configure("Results.TLabelframe.Label", background=self.result_bg_color, font=("Segoe UI", 12, "bold"), foreground=self.header_color)


    def _create_widgets(self):
        """Uygulamanın tüm ana bileşenlerini (widget) oluşturur ve verilerle doldurur."""
        main_frame = ttk.Frame(self.root, padding=10)
        main_frame.pack(expand=True, fill="both")

        self.tab_control = ttk.Notebook(main_frame)
        self.tab_control.pack(expand=True, fill="both")

        tabs = {
            "Maliyet & Kar Hesaplama": self._create_cost_profit_tab,
            "Filament Ayarları": self._create_filament_tab,
            "Pazaryeri Ayarları": self._create_marketplace_tab,
            "Kargo Ücretleri": self._create_shipping_tab,
            "Genel Ayarlar": self._create_settings_tab,
        }
        for text, creation_method in tabs.items():
            tab_frame = ttk.Frame(self.tab_control, padding=10)
            self.tab_control.add(tab_frame, text=text)
            creation_method(tab_frame)

        self._populate_all_widgets()

        status_bar = ttk.Label(self.root, textvariable=self.status_var, style="Status.TLabel", anchor='w')
        status_bar.pack(side="bottom", fill="x")

    def _populate_all_widgets(self):
        self._refresh_filament_list()
        self._update_filament_combobox()
        self._refresh_shipping_list()
        self._update_all_shipping_comboboxes()
        self._refresh_marketplace_list()
        self._update_all_marketplace_comboboxes()
        self._clear_marketplace_fields()
        default_power = self.settings.get("default_device_power", 130.0)
        self.cost_entries["device_power"].delete(0, tk.END)
        self.cost_entries["device_power"].insert(0, str(default_power))

    def _load_data(self):
        if os.path.exists(DATA_FILE_NAME):
            try:
                with open(DATA_FILE_NAME, "r", encoding='utf-8') as f:
                    data = json.load(f)
                    self.filaments = data.get(KEY_FILAMENTS, [])
                    self.electricity_prices.update(data.get(KEY_ELECTRICITY, {}))
                    self.shipping_prices = data.get(KEY_SHIPPING, {})
                    self.marketplaces = data.get(KEY_MARKETPLACES, DEFAULT_MARKETPLACES.copy())
                    self.settings.update(data.get(KEY_SETTINGS, {}))
            except (json.JSONDecodeError, TypeError):
                messagebox.showwarning("Yükleme Hatası", f"{DATA_FILE_NAME} dosyası bozuk. Varsayılanlar kullanılacak.")
                self.marketplaces = DEFAULT_MARKETPLACES.copy()
        else:
            self._save_data()

    def _save_data(self):
        data_to_save = {
            KEY_FILAMENTS: self.filaments,
            KEY_ELECTRICITY: self.electricity_prices,
            KEY_SHIPPING: self.shipping_prices,
            KEY_MARKETPLACES: self.marketplaces,
            KEY_SETTINGS: self.settings,
        }
        try:
            with open(DATA_FILE_NAME, "w", encoding='utf-8') as f:
                json.dump(data_to_save, f, indent=4, ensure_ascii=False)
            self.update_status("Değişiklikler başarıyla kaydedildi.")
        except Exception as e:
            messagebox.showerror("Kaydetme Hatası", f"Veri kaydedilirken bir hata oluştu: {e}")

    def _on_close(self):
        self._save_data()
        self.root.destroy()

    def _validate_numeric_input(self, P: str) -> bool:
        if P == "" or P == ".": return True
        try:
            return float(P) >= 0
        except ValueError:
            return False

    def update_status(self, message: str):
        self.status_var.set(f" {message}")
        
    def _create_treeview(self, parent, columns, headings) -> ttk.Treeview:
        frame = ttk.Frame(parent)
        frame.pack(expand=True, fill='both')
        tree = ttk.Treeview(frame, columns=columns, show="headings")
        for col, heading in zip(columns, headings):
            tree.heading(col, text=heading)
        tree.grid(row=0, column=0, sticky='nsew')
        scrollbar = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=scrollbar.set)
        scrollbar.grid(row=0, column=1, sticky='ns')
        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)
        return tree

    def _create_filament_tab(self, parent: ttk.Frame):
        parent.columnconfigure(0, weight=1)
        parent.rowconfigure(1, weight=1)
        input_frame = ttk.LabelFrame(parent, text="Yeni Filament Ekle/Sil")
        input_frame.grid(row=0, column=0, sticky="ew", padx=5, pady=5)
        input_frame.columnconfigure(1, weight=1)
        ttk.Label(input_frame, text="Filament Türü:").grid(row=0, column=0, padx=5, pady=5, sticky="w")
        self.filament_type_entry = ttk.Entry(input_frame, width=30)
        self.filament_type_entry.grid(row=0, column=1, padx=5, pady=5, sticky="ew")
        ttk.Label(input_frame, text="Fiyat (TL/kg):").grid(row=1, column=0, padx=5, pady=5, sticky="w")
        vcmd_numeric = (self.root.register(self._validate_numeric_input), '%P')
        self.filament_price_entry = ttk.Entry(input_frame, validate='key', validatecommand=vcmd_numeric, width=15)
        self.filament_price_entry.grid(row=1, column=1, padx=5, pady=5, sticky="w")
        btn_frame = ttk.Frame(input_frame)
        btn_frame.grid(row=2, column=0, columnspan=2, pady=10)
        ttk.Button(btn_frame, text="Ekle", command=self._add_filament).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="Seçileni Sil", command=self._delete_filament).pack(side=tk.LEFT, padx=5)
        list_frame = ttk.LabelFrame(parent, text="Kayıtlı Filamentler")
        list_frame.grid(row=1, column=0, sticky="nsew", padx=5, pady=5)
        self.filament_tree = self._create_treeview(list_frame, ("Tür", "Fiyat"), ("Filament Türü", "Fiyat (TL/kg)"))
        self.filament_tree.column("Fiyat", width=120, anchor="e", stretch=tk.NO)

    def _create_cost_profit_tab(self, parent: ttk.Frame):
        parent.columnconfigure(0, weight=1)
        parent.columnconfigure(1, weight=2)
        parent.rowconfigure(0, weight=1)
        left_panel = ttk.Frame(parent)
        left_panel.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        right_panel = ttk.Frame(parent, style="Results.TFrame", padding=15)
        right_panel.grid(row=0, column=1, sticky="nsew")
        right_panel.columnconfigure(0, weight=1)
        right_panel.rowconfigure(0, weight=1)
        vcmd_numeric = (self.root.register(self._validate_numeric_input), '%P')
        input_frame = ttk.LabelFrame(left_panel, text="Baskı Bilgileri")
        input_frame.pack(fill="x", expand=False, pady=(0, 10))
        input_frame.columnconfigure(1, weight=1)
        self.cost_entries = {}
        input_fields = {
            "filament_amount": "Filament Miktarı (g):",
            "print_time": "Baskı Süresi (saat):",
            "device_power": "Cihaz Gücü (W):",
        }
        for i, (key, text) in enumerate(input_fields.items()):
            ttk.Label(input_frame, text=text).grid(row=i, column=0, padx=5, pady=5, sticky="w")
            entry = ttk.Entry(input_frame, validate='key', validatecommand=vcmd_numeric)
            entry.grid(row=i, column=1, padx=5, pady=5, sticky="ew")
            self.cost_entries[key] = entry
        ttk.Label(input_frame, text="Filament Türü:").grid(row=len(input_fields), column=0, padx=5, pady=5, sticky="w")
        self.calc_filament_combo = ttk.Combobox(input_frame, state="readonly")
        self.calc_filament_combo.grid(row=len(input_fields), column=1, padx=5, pady=5, sticky="ew")
        extra_frame = ttk.LabelFrame(left_panel, text="Ek Maliyetler, Kar ve Pazaryeri")
        extra_frame.pack(fill="x", expand=False)
        extra_frame.columnconfigure(1, weight=1)
        ttk.Label(extra_frame, text="İstenen Kar Oranı (%):").grid(row=0, column=0, padx=5, pady=5, sticky="w")
        self.profit_margin_entry = ttk.Entry(extra_frame, validate='key', validatecommand=vcmd_numeric)
        self.profit_margin_entry.grid(row=0, column=1, padx=5, pady=5, sticky="ew")
        self.profit_margin_entry.insert(0, "100.0")
        ttk.Label(extra_frame, text="Diğer Giderler (TL):").grid(row=1, column=0, padx=5, pady=5, sticky="w")
        self.other_costs_entry = ttk.Entry(extra_frame, validate='key', validatecommand=vcmd_numeric)
        self.other_costs_entry.grid(row=1, column=1, padx=5, pady=5, sticky="ew")
        self.other_costs_entry.insert(0, "0.0")
        ttk.Label(extra_frame, text="Kargo Firması:").grid(row=2, column=0, padx=5, pady=5, sticky="w")
        self.calc_shipping_company_combo = ttk.Combobox(extra_frame, state="readonly")
        self.calc_shipping_company_combo.grid(row=2, column=1, padx=5, pady=5, sticky="ew")
        self.calc_shipping_company_combo.bind("<<ComboboxSelected>>", self._update_shipping_desi_combobox)
        ttk.Label(extra_frame, text="Desi/Ağırlık:").grid(row=3, column=0, padx=5, pady=5, sticky="w")
        self.calc_shipping_desi_combo = ttk.Combobox(extra_frame, state="readonly")
        self.calc_shipping_desi_combo.grid(row=3, column=1, padx=5, pady=5, sticky="ew")
        ttk.Label(extra_frame, text="Satış Pazaryeri:").grid(row=4, column=0, padx=5, pady=5, sticky="w")
        self.calc_marketplace_combo = ttk.Combobox(extra_frame, state="readonly")
        self.calc_marketplace_combo.grid(row=4, column=1, padx=5, pady=5, sticky="ew")
        ttk.Button(left_panel, text="Hesapla", command=self._calculate_and_display, style="TButton").pack(pady=20, fill='x', ipady=5)
        results_frame = ttk.LabelFrame(right_panel, text="Hesaplama Sonuçları", style="Results.TLabelframe")
        results_frame.pack(fill="both", expand=True)
        results_frame.columnconfigure(1, weight=1)
        self.result_labels: Dict[str, ttk.Label] = {}
        result_fields = {
            "filament_cost": "Filament Maliyeti:", "elec_cost": "Elektrik Maliyeti:",
            "shipping_cost": "Kargo Maliyeti:", "other_costs": "Diğer Giderler:",
            "sep1": None, "total_base_cost": "Toplam Temel Maliyet:", "sep2": None,
            "commission": "Platform Komisyonu:", "total_final_cost": "Komisyonlu Toplam Maliyet:",
            "sep3": None, "custom_sale_price": "Satış Fiyatı:", "net_profit": "Net Kar:",
        }
        row_counter = 0
        for key, text in result_fields.items():
            if text is None:
                ttk.Separator(results_frame, orient='horizontal').grid(row=row_counter, column=0, columnspan=2, sticky='ew', pady=8)
            else:
                label_widget = ttk.Label(results_frame, text=text, anchor="w", style="Results.TLabel", font=("Segoe UI", 10))
                label_widget.grid(row=row_counter, column=0, padx=5, pady=6, sticky="w")
                self.result_labels[f"{key}_text"] = label_widget
                value_label = ttk.Label(results_frame, text="0.00 TL", anchor="e", style="Results.TLabel", font=("Segoe UI", 11, "bold"))
                value_label.grid(row=row_counter, column=1, padx=5, pady=6, sticky="ew")
                self.result_labels[key] = value_label
            row_counter += 1
        ttk.Separator(results_frame, orient='horizontal').grid(row=row_counter, column=0, columnspan=2, sticky='ew', pady=15)
        row_counter += 1
        add_to_sales_btn = ttk.Button(results_frame, text="Bu Satışı Excel'e Ekle", command=self._add_to_sales)
        add_to_sales_btn.grid(row=row_counter, column=0, columnspan=2, sticky='ew', ipady=4)
        # DEĞİŞTİRİLDİ: ToolTip kaldırıldı.
        # ToolTip(add_to_sales_btn, f"Mevcut hesaplama sonucunu '{EXCEL_FILE_NAME}' dosyasına kaydeder.")

    def _create_marketplace_tab(self, parent: ttk.Frame):
        parent.columnconfigure(0, weight=1)
        parent.rowconfigure(1, weight=1)
        vcmd_numeric = (self.root.register(self._validate_numeric_input), '%P')
        input_frame = ttk.LabelFrame(parent, text="Pazaryeri Ekle/Düzenle")
        input_frame.grid(row=0, column=0, sticky="ew", padx=5, pady=5)
        input_frame.columnconfigure(1, weight=1)
        ttk.Label(input_frame, text="Pazaryeri Adı:").grid(row=0, column=0, padx=5, pady=5, sticky="w")
        self.mplace_name_entry = ttk.Entry(input_frame)
        self.mplace_name_entry.grid(row=0, column=1, padx=5, pady=5, sticky="ew")
        ttk.Label(input_frame, text="Komisyon Oranı (%):").grid(row=1, column=0, padx=5, pady=5, sticky="w")
        self.mplace_rate_entry = ttk.Entry(input_frame, validate='key', validatecommand=vcmd_numeric, width=15)
        self.mplace_rate_entry.grid(row=1, column=1, padx=5, pady=5, sticky="w")
        ttk.Label(input_frame, text="Sabit Komisyon (TL):").grid(row=2, column=0, padx=5, pady=5, sticky="w")
        self.mplace_fixed_entry = ttk.Entry(input_frame, validate='key', validatecommand=vcmd_numeric, width=15)
        self.mplace_fixed_entry.grid(row=2, column=1, padx=5, pady=5, sticky="w")
        btn_frame = ttk.Frame(input_frame)
        btn_frame.grid(row=3, column=0, columnspan=2, pady=10)
        self.mplace_save_btn = ttk.Button(btn_frame, text="Ekle", command=self._save_or_update_marketplace)
        self.mplace_save_btn.pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="Seçileni Sil", command=self._delete_marketplace).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="Temizle/Yeni", command=self._clear_marketplace_fields).pack(side=tk.LEFT, padx=5)
        list_frame = ttk.LabelFrame(parent, text="Kayıtlı Pazaryerleri")
        list_frame.grid(row=1, column=0, sticky="nsew", padx=5, pady=5)
        self.mplace_tree = self._create_treeview(list_frame, ("Ad", "Oran", "Sabit"), ("Pazaryeri Adı", "Komisyon (%)", "Sabit Komisyon (TL)"))
        self.mplace_tree.column("Oran", width=120, anchor="e", stretch=tk.NO)
        self.mplace_tree.column("Sabit", width=140, anchor="e", stretch=tk.NO)
        self.mplace_tree.bind("<<TreeviewSelect>>", self._load_marketplace_for_edit)

    def _create_shipping_tab(self, parent: ttk.Frame):
        parent.columnconfigure(0, weight=1)
        parent.rowconfigure(1, weight=1)
        vcmd_numeric = (self.root.register(self._validate_numeric_input), '%P')
        input_frame = ttk.LabelFrame(parent, text="Yeni Kargo Fiyatı Ekle/Sil")
        input_frame.grid(row=0, column=0, sticky="ew", padx=5, pady=5)
        input_frame.columnconfigure(1, weight=1)
        ttk.Label(input_frame, text="Kargo Firması:").grid(row=0, column=0, padx=5, pady=5, sticky="w")
        self.ship_company_combo = ttk.Combobox(input_frame)
        self.ship_company_combo.grid(row=0, column=1, padx=5, pady=5, sticky="ew")
        ttk.Label(input_frame, text="Desi/Ağırlık:").grid(row=1, column=0, padx=5, pady=5, sticky="w")
        self.ship_desi_entry = ttk.Entry(input_frame, width=15)
        self.ship_desi_entry.grid(row=1, column=1, padx=5, pady=5, sticky="w")
        ttk.Label(input_frame, text="Fiyat (TL):").grid(row=2, column=0, padx=5, pady=5, sticky="w")
        self.ship_price_entry = ttk.Entry(input_frame, validate='key', validatecommand=vcmd_numeric, width=15)
        self.ship_price_entry.grid(row=2, column=1, padx=5, pady=5, sticky="w")
        btn_frame = ttk.Frame(input_frame)
        btn_frame.grid(row=3, column=0, columnspan=2, pady=10)
        ttk.Button(btn_frame, text="Ekle/Güncelle", command=self._add_shipping_price).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="Seçileni Sil", command=self._delete_shipping_price).pack(side=tk.LEFT, padx=5)
        list_frame = ttk.LabelFrame(parent, text="Kayıtlı Kargo Fiyatları")
        list_frame.grid(row=1, column=0, sticky="nsew", padx=5, pady=5)
        self.shipping_tree = self._create_treeview(list_frame, ("Firma", "Desi", "Fiyat"), ("Kargo Firması", "Desi/Ağırlık", "Fiyat (TL)"))
        self.shipping_tree.column("Desi", width=120, anchor="center", stretch=tk.NO)
        self.shipping_tree.column("Fiyat", width=120, anchor="e", stretch=tk.NO)

    def _create_settings_tab(self, parent: ttk.Frame):
        parent.columnconfigure(0, weight=1)
        vcmd_numeric = (self.root.register(self._validate_numeric_input), '%P')
        elec_frame = ttk.LabelFrame(parent, text="Elektrik Fiyatları (TL/kWh)")
        elec_frame.grid(row=0, column=0, sticky="ew", padx=5, pady=5)
        elec_frame.columnconfigure(1, weight=1)
        ttk.Label(elec_frame, text="Ticari Tarife:").grid(row=0, column=0, padx=5, pady=5, sticky="w")
        self.comm_elec_entry = ttk.Entry(elec_frame, validate='key', validatecommand=vcmd_numeric)
        self.comm_elec_entry.grid(row=0, column=1, padx=5, pady=5, sticky="ew")
        self.comm_elec_entry.insert(0, str(self.electricity_prices.get("ticari", 0.0)))
        ttk.Label(elec_frame, text="Mesken Tarife:").grid(row=1, column=0, padx=5, pady=5, sticky="w")
        self.res_elec_entry = ttk.Entry(elec_frame, validate='key', validatecommand=vcmd_numeric)
        self.res_elec_entry.grid(row=1, column=1, padx=5, pady=5, sticky="ew")
        self.res_elec_entry.insert(0, str(self.electricity_prices.get("mesken", 0.0)))
        costing_frame = ttk.LabelFrame(parent, text="Genel Hesaplama Ayarları")
        costing_frame.grid(row=1, column=0, sticky="ew", padx=5, pady=15)
        costing_frame.columnconfigure(1, weight=1)
        ttk.Label(costing_frame, text="Varsayılan Cihaz Gücü (W):").grid(row=0, column=0, padx=5, pady=5, sticky="w")
        self.default_power_entry = ttk.Entry(costing_frame, validate='key', validatecommand=vcmd_numeric)
        self.default_power_entry.grid(row=0, column=1, padx=5, pady=5, sticky="ew")
        self.default_power_entry.insert(0, str(self.settings.get("default_device_power", 130.0)))
        ttk.Label(costing_frame, text="Varsayılan Elektrik Tarifesi:").grid(row=1, column=0, padx=5, pady=5, sticky="w")
        self.costing_elec_var = tk.StringVar(value=self.settings.get("costing_electricity_type", "ticari"))
        rb_frame = ttk.Frame(costing_frame)
        rb_frame.grid(row=1, column=1, padx=5, pady=5, sticky="w")
        ttk.Radiobutton(rb_frame, text="Ticari", variable=self.costing_elec_var, value="ticari").pack(side=tk.LEFT, padx=5)
        ttk.Radiobutton(rb_frame, text="Mesken", variable=self.costing_elec_var, value="mesken").pack(side=tk.LEFT, padx=5)
        ttk.Button(parent, text="Ayarları Kaydet", command=self._update_settings).grid(row=2, column=0, pady=20)
        support_frame = ttk.LabelFrame(parent, text="Destek & İletişim")
        support_frame.grid(row=3, column=0, sticky="ew", padx=5, pady=15)
        support_frame.columnconfigure(0, weight=1)
        discord_btn = ttk.Button(support_frame, text="Geliştiriciye Discord'dan Ulaş", command=self._open_discord_support)
        discord_btn.grid(row=0, column=0, sticky="ew", padx=5, pady=5)
        # DEĞİŞTİRİLDİ: ToolTip kaldırıldı.
        # ToolTip(discord_btn, f"Discord uygulamasını açarak geliştiriciye (ID: {DISCORD_USER_ID}) mesaj göndermenizi sağlar.")

    def _add_filament(self):
        f_type = self.filament_type_entry.get().strip()
        f_price_str = self.filament_price_entry.get()
        if not all([f_type, f_price_str]):
            messagebox.showwarning("Eksik Bilgi", "Filament türü ve fiyatı boş bırakılamaz.")
            return
        try:
            f_price = float(f_price_str)
            if any(f[0].lower() == f_type.lower() for f in self.filaments):
                messagebox.showwarning("Mevcut Filament", f"'{f_type}' zaten listede mevcut.")
                return
            self.filaments.append((f_type, f_price))
            self.filaments.sort(key=lambda x: x[0].lower())
            self._refresh_filament_list()
            self._update_filament_combobox()
            self._save_data()
            self.filament_type_entry.delete(0, tk.END)
            self.filament_price_entry.delete(0, tk.END)
            self.update_status(f"'{f_type}' filamenti eklendi.")
        except ValueError:
            messagebox.showerror("Hata", "Fiyat geçersiz! Lütfen sayısal bir değer girin.")

    def _delete_filament(self):
        selected_items = self.filament_tree.selection()
        if not selected_items:
            messagebox.showwarning("Seçim Yapılmadı", "Lütfen silmek için bir filament seçin.")
            return
        item_values = self.filament_tree.item(selected_items[0], "values")
        name_to_delete = item_values[0]
        if messagebox.askyesno("Onay", f"'{name_to_delete}' filamentini silmek istediğinizden emin misiniz?"):
            self.filaments = [f for f in self.filaments if f[0] != name_to_delete]
            self._refresh_filament_list()
            self._update_filament_combobox()
            self._save_data()
            self.update_status(f"'{name_to_delete}' filamenti silindi.")

    def _refresh_filament_list(self):
        self.filament_tree.delete(*self.filament_tree.get_children())
        for f_type, f_price in self.filaments:
            self.filament_tree.insert("", tk.END, values=(f_type, f"{f_price:.2f}"))

    def _update_filament_combobox(self):
        types = sorted([f[0] for f in self.filaments], key=str.lower)
        self.calc_filament_combo["values"] = types
        if types: self.calc_filament_combo.set(types[0])
        else: self.calc_filament_combo.set("")

    def _save_or_update_marketplace(self):
        name = self.mplace_name_entry.get().strip()
        rate_str = self.mplace_rate_entry.get()
        fixed_str = self.mplace_fixed_entry.get()
        if not all([name, rate_str, fixed_str]):
            messagebox.showwarning("Eksik Bilgi", "Tüm pazaryeri alanları doldurulmalıdır.")
            return
        try:
            rate = float(rate_str)
            fixed = float(fixed_str)
            if self.selected_marketplace_name:
                original_name = self.selected_marketplace_name
                if name != original_name and name.lower() in (k.lower() for k in self.marketplaces):
                    messagebox.showwarning("Mevcut İsim", f"'{name}' adında başka bir pazaryeri zaten var.")
                    return
                if original_name in self.marketplaces:
                    del self.marketplaces[original_name]
                action = "güncellendi"
            else:
                if name.lower() in (k.lower() for k in self.marketplaces):
                    messagebox.showwarning("Mevcut Pazaryeri", f"'{name}' adında bir pazaryeri zaten mevcut.")
                    return
                action = "eklendi"
            self.marketplaces[name] = {"rate": rate, "fixed": fixed}
            self._save_data()
            self._refresh_marketplace_list()
            self._update_all_marketplace_comboboxes()
            self._clear_marketplace_fields()
            self.update_status(f"Pazaryeri '{name}' başarıyla {action}.")
        except ValueError:
            messagebox.showerror("Hata", "Oran ve Sabit Komisyon için geçerli sayısal değerler girin.")

    def _delete_marketplace(self):
        if not self.selected_marketplace_name:
            messagebox.showwarning("Seçim Yapılmadı", "Lütfen silmek için bir pazaryeri seçin.")
            return
        name_to_delete = self.selected_marketplace_name
        if messagebox.askyesno("Onay", f"'{name_to_delete}' pazaryerini silmek istediğinizden emin misiniz?"):
            if name_to_delete in self.marketplaces:
                del self.marketplaces[name_to_delete]
                self._save_data()
                self._refresh_marketplace_list()
                self._update_all_marketplace_comboboxes()
                self._clear_marketplace_fields()
                self.update_status(f"'{name_to_delete}' pazaryeri silindi.")

    def _load_marketplace_for_edit(self, event=None):
        selected_items = self.mplace_tree.selection()
        if not selected_items: return
        name = self.mplace_tree.item(selected_items[0], "values")[0]
        if name in self.marketplaces:
            self.selected_marketplace_name = name
            data = self.marketplaces[name]
            self.mplace_name_entry.delete(0, tk.END); self.mplace_name_entry.insert(0, name)
            self.mplace_rate_entry.delete(0, tk.END); self.mplace_rate_entry.insert(0, f"{data.get('rate', 0.0):.2f}")
            self.mplace_fixed_entry.delete(0, tk.END); self.mplace_fixed_entry.insert(0, f"{data.get('fixed', 0.0):.2f}")
            self.mplace_save_btn.config(text="Güncelle")

    def _clear_marketplace_fields(self):
        self.selected_marketplace_name = None
        self.mplace_name_entry.delete(0, tk.END)
        self.mplace_rate_entry.delete(0, tk.END)
        self.mplace_fixed_entry.delete(0, tk.END)
        self.mplace_save_btn.config(text="Ekle")
        if self.mplace_tree.selection():
            self.mplace_tree.selection_remove(self.mplace_tree.selection())

    def _refresh_marketplace_list(self):
        self.mplace_tree.delete(*self.mplace_tree.get_children())
        for name in sorted(self.marketplaces.keys(), key=str.lower):
            data = self.marketplaces[name]
            self.mplace_tree.insert("", tk.END, values=(name, f"{data.get('rate', 0.0):.2f}", f"{data.get('fixed', 0.0):.2f}"))

    def _update_all_marketplace_comboboxes(self):
        names = sorted(list(self.marketplaces.keys()), key=str.lower)
        self.calc_marketplace_combo["values"] = names
        default = "Diğer/Yok" if "Diğer/Yok" in names else (names[0] if names else "")
        self.calc_marketplace_combo.set(default)

    def _add_shipping_price(self):
        company = self.ship_company_combo.get().strip()
        desi = self.ship_desi_entry.get().strip()
        price_str = self.ship_price_entry.get()
        if not all([company, desi, price_str]):
            messagebox.showwarning("Eksik Bilgi", "Firma, desi ve fiyat alanları doldurulmalıdır.")
            return
        try:
            price = float(price_str)
            if company not in self.shipping_prices:
                self.shipping_prices[company] = []
            existing_desi = next((item for item in self.shipping_prices[company] if item[0] == desi), None)
            if existing_desi:
                self.shipping_prices[company].remove(existing_desi)
            self.shipping_prices[company].append((desi, price))
            self.shipping_prices[company].sort(key=lambda item: (0, float(item[0])) if item[0].replace('.', '', 1).isdigit() else (1, item[0]))
            self._save_data()
            self._refresh_shipping_list()
            self._update_all_shipping_comboboxes()
            self.ship_desi_entry.delete(0, tk.END)
            self.ship_price_entry.delete(0, tk.END)
            self.update_status(f"'{company}' için kargo fiyatı eklendi/güncellendi.")
        except ValueError:
            messagebox.showerror("Hata", "Fiyat geçersiz! Lütfen sayısal bir değer girin.")

    def _delete_shipping_price(self):
        selected_items = self.shipping_tree.selection()
        if not selected_items:
            messagebox.showwarning("Seçim Yapılmadı", "Lütfen silmek için bir kargo fiyatı seçin.")
            return
        company, desi, _ = self.shipping_tree.item(selected_items[0], "values")
        if messagebox.askyesno("Onay", f"'{company}' firmasından '{desi}' kaydını silmek istediğinizden emin misiniz?"):
            if company in self.shipping_prices:
                self.shipping_prices[company] = [item for item in self.shipping_prices[company] if item[0] != desi]
                if not self.shipping_prices[company]:
                    del self.shipping_prices[company]
                self._save_data()
                self._refresh_shipping_list()
                self._update_all_shipping_comboboxes()
                self.update_status(f"Kargo kaydı ({company} - {desi}) silindi.")

    def _refresh_shipping_list(self):
        self.shipping_tree.delete(*self.shipping_tree.get_children())
        for company in sorted(self.shipping_prices.keys(), key=str.lower):
            for desi, price in self.shipping_prices.get(company, []):
                self.shipping_tree.insert("", tk.END, values=(company, desi, f"{price:.2f}"))

    def _update_all_shipping_comboboxes(self):
        companies = sorted(list(self.shipping_prices.keys()), key=str.lower)
        self.ship_company_combo['values'] = companies
        self.calc_shipping_company_combo['values'] = companies
        if companies:
            self.ship_company_combo.set(companies[0])
            self.calc_shipping_company_combo.set(companies[0])
        else:
            self.ship_company_combo.set("")
            self.calc_shipping_company_combo.set("")
        self._update_shipping_desi_combobox()

    def _update_shipping_desi_combobox(self, event=None):
        company = self.calc_shipping_company_combo.get()
        desi_options = [item[0] for item in self.shipping_prices.get(company, [])]
        self.calc_shipping_desi_combo["values"] = desi_options
        if desi_options: self.calc_shipping_desi_combo.set(desi_options[0])
        else: self.calc_shipping_desi_combo.set("")

    def _update_settings(self):
        try:
            comm_price = float(self.comm_elec_entry.get())
            res_price = float(self.res_elec_entry.get())
            power = float(self.default_power_entry.get())
            costing_type = self.costing_elec_var.get()
            self.electricity_prices["ticari"] = comm_price
            self.electricity_prices["mesken"] = res_price
            self.settings["default_device_power"] = power
            self.settings["costing_electricity_type"] = costing_type
            self._save_data()
            self.cost_entries["device_power"].delete(0, tk.END)
            self.cost_entries["device_power"].insert(0, str(power))
            messagebox.showinfo("Başarılı", "Genel ayarlar güncellendi.")
        except (ValueError, TypeError):
            messagebox.showerror("Hata", "Geçersiz giriş! Lütfen sayısal alanları doldurun.")
            
    def _open_discord_support(self):
        url = f"discord://-/users/{DISCORD_USER_ID}"
        try:
            webbrowser.open(url)
            self.update_status("Discord uygulaması açılıyor...")
        except Exception as e:
            messagebox.showerror("Hata", f"Discord açılamadı. Yüklü olduğundan emin olun.\n\nHata: {e}")

    def _calculate_and_display(self):
        self.last_calculation_results = None
        self.last_calculation_inputs = None
        try:
            inputs = {key: float(entry.get() or 0) for key, entry in self.cost_entries.items()}
            inputs["other_costs"] = float(self.other_costs_entry.get() or 0)
            inputs["profit_margin"] = float(self.profit_margin_entry.get() or 0)
            inputs["filament_type"] = self.calc_filament_combo.get()
            inputs["shipping_company"] = self.calc_shipping_company_combo.get()
            inputs["shipping_desi"] = self.calc_shipping_desi_combo.get()
            inputs["marketplace"] = self.calc_marketplace_combo.get()
            if not all([inputs["filament_type"], inputs["marketplace"]]):
                messagebox.showwarning("Eksik Bilgi", "Filament türü ve pazaryeri seçin.")
                return
        except (ValueError, TypeError):
            messagebox.showerror("Geçersiz Giriş", "Tüm sayısal alanları doldurun.")
            return
        try:
            results = self._perform_calculations(inputs)
            self.last_calculation_results = results
            self.last_calculation_inputs = inputs
        except Exception as e:
            messagebox.showerror("Hesaplama Hatası", f"Hesaplama sırasında hata: {e}")
            return
        self._update_results_display(results, inputs)
        self.update_status("Hesaplama tamamlandı.")

    def _perform_calculations(self, inputs: Dict[str, Any]) -> Dict[str, float]:
        results = {}
        f_price_kg = next((price for type_, price in self.filaments if type_ == inputs["filament_type"]), 0.0)
        results["filament_cost"] = (inputs["filament_amount"] / 1000.0) * f_price_kg
        elec_type = self.settings.get("costing_electricity_type", "ticari")
        elec_price = self.electricity_prices.get(elec_type, 0.0)
        kwh = (inputs["device_power"] * inputs["print_time"]) / 1000.0
        results["elec_cost"] = kwh * elec_price
        shipping_cost = 0.0
        if inputs["shipping_company"] and inputs["shipping_desi"]:
            shipping_cost = next((price for d, price in self.shipping_prices.get(inputs["shipping_company"], []) if d == inputs["shipping_desi"]), 0.0)
        results["shipping_cost"] = shipping_cost
        base_cost = results["filament_cost"] + results["elec_cost"] + results["shipping_cost"] + inputs["other_costs"]
        results["total_base_cost"] = base_cost
        mplace_info = self.marketplaces.get(inputs["marketplace"], {"rate": 0.0, "fixed": 0.0})
        rate = mplace_info.get("rate", 0.0) / 100.0
        fixed = mplace_info.get("fixed", 0.0)
        profit_margin = inputs["profit_margin"] / 100.0
        denominator = 1 - rate
        if denominator <= 0:
            sale_price, commission, total_final_cost, net_profit = [float('inf')] * 4
        else:
            sale_price = (base_cost * (1 + profit_margin) + fixed) / denominator
            commission = (sale_price * rate) + fixed
            total_final_cost = base_cost + commission
            net_profit = sale_price - total_final_cost
        results.update({
            "custom_sale_price": sale_price, "commission": commission,
            "total_final_cost": total_final_cost, "net_profit": net_profit
        })
        return results

    def _update_results_display(self, results: Dict[str, float], inputs: Dict[str, Any]):
        self.result_labels["filament_cost"].config(text=f"{results.get('filament_cost', 0):.2f} TL")
        self.result_labels["elec_cost"].config(text=f"{results.get('elec_cost', 0):.2f} TL")
        self.result_labels["shipping_cost"].config(text=f"{results.get('shipping_cost', 0):.2f} TL")
        self.result_labels["other_costs"].config(text=f"{inputs.get('other_costs', 0):.2f} TL")
        self.result_labels["total_base_cost"].config(text=f"{results.get('total_base_cost', 0):.2f} TL")
        elec_type_tr = self.settings.get("costing_electricity_type", "ticari").capitalize()
        self.result_labels["elec_cost_text"].config(text=f"Elektrik Maliyeti ({elec_type_tr}):")
        self.result_labels["commission_text"].config(text=f"Komisyon ({inputs.get('marketplace')}):")
        self.result_labels["custom_sale_price_text"].config(text=f"Satış Fiyatı (%{inputs.get('profit_margin', 100):.1f} Kar ile):")
        for key in ["commission", "total_final_cost", "custom_sale_price", "net_profit"]:
            value = results.get(key, 0.0)
            text = "Hesaplanamadı" if value == float('inf') else f"{value:.2f} TL"
            self.result_labels[key].config(text=text)

    def _add_to_sales(self):
        if not self.last_calculation_results or not self.last_calculation_inputs:
            messagebox.showwarning("Hesaplama Gerekli", "Lütfen önce bir maliyet hesaplaması yapın.")
            return

        dialog = CustomAskStringDialog(self.root, "Ürün Adı Girin", "Lütfen satılan ürünün adını girin:")
        product_name = dialog.result

        if not product_name:
            self.update_status("Satış ekleme işlemi iptal edildi.")
            return
            
        try:
            # Gerekli verileri topla
            filament_type = self.last_calculation_inputs.get("filament_type", "Bilinmiyor")
            grams = self.last_calculation_inputs.get("filament_amount", 0.0)
            cost = self.last_calculation_results.get("total_final_cost", 0.0)
            sale_price = self.last_calculation_results.get("custom_sale_price", 0.0)
            profit = self.last_calculation_results.get("net_profit", 0.0)
            headers = ["Ürün Adı", "Filament Türü", "GR", "Maliyet (TL)", "Satış (TL)", "Kar (TL)"]
            
            # Çalışma kitabını yükle veya oluştur
            if not os.path.exists(EXCEL_FILE_NAME):
                workbook = openpyxl.Workbook()
                # Varsayılan sayfayı "Satışlar" olarak yeniden adlandır
                sales_sheet = workbook.active
                sales_sheet.title = "Satışlar"
                sales_sheet.append(headers)
                # İlk defa oluşturulduğunda boş bir Özet sayfası da ekle
                workbook.create_sheet("Özet")
            else:
                workbook = openpyxl.load_workbook(EXCEL_FILE_NAME)
                # "Satışlar" sayfasını al veya oluştur
                if "Satışlar" in workbook.sheetnames:
                    sales_sheet = workbook["Satışlar"]
                else: # Eski dosya formatıyla uyumluluk için
                    sales_sheet = workbook.create_sheet("Satışlar")
                    sales_sheet.append(headers)

            # Yeni satış verisini ekle
            data_row = [product_name, filament_type, grams, round(cost, 2), round(sale_price, 2), round(profit, 2)]
            sales_sheet.append(data_row)
            
            # Özeti aynı çalışma kitabı nesnesi üzerinde güncelle
            self._update_excel_summary(workbook)

            # Tüm değişiklikleri tek seferde kaydet
            workbook.save(EXCEL_FILE_NAME)

            messagebox.showinfo("Başarılı", f"Satış başarıyla '{EXCEL_FILE_NAME}' dosyasına eklendi ve özet güncellendi.")
            self.update_status(f"'{product_name}' satışı Excel'e kaydedildi.")

        except PermissionError:
             messagebox.showerror("İzin Hatası", f"'{EXCEL_FILE_NAME}' dosyasına yazılamıyor.\nLütfen dosyanın başka bir programda açık olmadığından emin olun.")
        except Exception as e:
            messagebox.showerror("Excel Yazma Hatası", f"Veri Excel'e yazılırken bir hata oluştu:\n\n{e}")
            traceback.print_exc()

    def _update_excel_summary(self, workbook):
        try:
            # Gerekli sayfaları al
            if "Satışlar" not in workbook.sheetnames:
                print("Satışlar sayfası bulunamadı. Özet oluşturulamadı.")
                return
            
            sales_sheet = workbook["Satışlar"]
            
            # "Özet" sayfasını temizleyip yeniden oluştur
            if "Özet" in workbook.sheetnames:
                del workbook["Özet"]
            summary_sheet = workbook.create_sheet("Özet")

            # Veri okuma ve hesaplama
            filament_totals = {}
            total_cost, total_sale, total_profit = 0.0, 0.0, 0.0
            
            headers = [cell.value for cell in sales_sheet[1]]
            try:
                filament_col = headers.index("Filament Türü")
                grams_col = headers.index("GR")
                cost_col = headers.index("Maliyet (TL)")
                sale_col = headers.index("Satış (TL)")
                profit_col = headers.index("Kar (TL)")
            except ValueError as e:
                print(f"Excel başlıklarında hata: '{e.args[0]}' sütunu bulunamadı. Özet oluşturulamadı.")
                return

            for row in sales_sheet.iter_rows(min_row=2, values_only=True):
                if not any(row): continue # Boş satırları atla
                try:
                    f_type = row[filament_col]
                    grams = float(row[grams_col] or 0)
                    cost = float(row[cost_col] or 0)
                    sale = float(row[sale_col] or 0)
                    profit = float(row[profit_col] or 0)

                    filament_totals[f_type] = filament_totals.get(f_type, 0) + grams
                    total_cost += cost
                    total_sale += sale
                    total_profit += profit
                except (TypeError, ValueError) as e:
                    print(f"Özet hesaplanırken bir satır atlandı (Hata: {e}): {row}")
                    continue

            # Özet verilerini "Özet" sayfasına yaz
            bold_font = Font(bold=True)
            
            summary_sheet.cell(row=1, column=1, value="Genel Toplamlar").font = bold_font
            summary_sheet.cell(row=2, column=1, value="Toplam Maliyet")
            summary_sheet.cell(row=2, column=2, value=total_cost).number_format = '#,##0.00 "TL"'
            summary_sheet.cell(row=3, column=1, value="Toplam Satış")
            summary_sheet.cell(row=3, column=2, value=total_sale).number_format = '#,##0.00 "TL"'
            summary_sheet.cell(row=4, column=1, value="Toplam Kar")
            summary_sheet.cell(row=4, column=2, value=total_profit).number_format = '#,##0.00 "TL"'
            
            current_row = 6
            summary_sheet.cell(row=current_row, column=1, value="Harcanan Filament").font = bold_font
            current_row += 1
            for f_type, total_grams in sorted(filament_totals.items()):
                summary_sheet.cell(row=current_row, column=1, value=f_type)
                summary_sheet.cell(row=current_row, column=2, value=total_grams).number_format = '#,##0.00 "gr"'
                current_row += 1

            # Sütun genişliklerini ayarla
            summary_sheet.column_dimensions['A'].width = 20
            summary_sheet.column_dimensions['B'].width = 20

        except Exception as e:
            print(f"Excel özeti güncellenirken bir hata oluştu: {e}")
            traceback.print_exc()

def main():
    try:
        root = tk.Tk()
        window_width, window_height = 1100, 800
        screen_width, screen_height = root.winfo_screenwidth(), root.winfo_screenheight()
        center_x = int(screen_width / 2 - window_width / 2)
        center_y = int(screen_height / 2 - window_height / 2)
        root.geometry(f'{window_width}x{window_height}+{center_x}+{center_y}')
        app = CostCalculatorApp(root)
        root.mainloop()
    except Exception as e:
        traceback.print_exc()
        messagebox.showerror(f"{APP_NAME} - Kritik Hata", f"Uygulama başlatılamadı:\n\n{e}")

if __name__ == "__main__":
    main()
